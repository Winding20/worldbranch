# -*- coding: utf-8 -*-
"""Verify the generated DOCX."""
import sys
from docx import Document

path = sys.argv[1]
d = Document(path)
paras = d.paragraphs
heads = [(p.style.name, p.text.strip()) for p in paras if p.style.name.startswith("Heading") and p.text.strip()]
print(f"file: {path}")
print(f"paragraphs={len(paras)} tables={len(d.tables)} headings={len(heads)}")
print("--- heading outline ---")
for s, t in heads:
    lvl = s.replace("Heading ", "")
    indent = "  " * (int(lvl) - 1)
    print(f"{indent}[H{lvl}] {t}")
print("--- table shapes ---")
for n, t in enumerate(d.tables):
    print(f"table{n}: rows={len(t.rows)} cols={len(t.columns)} first={t.rows[0].cells[0].text[:18]}")
print("--- empty-body check ---")
empty = [i for i, p in enumerate(paras) if not p.text.strip() and not p.style.name.startswith("Heading")]
print(f"empty paragraphs: {len(empty)}")
print("last paragraph:", repr(paras[-1].text[:60]))
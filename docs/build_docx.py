# -*- coding: utf-8 -*-
"""Generate WorldBranch project design DOCX from the markdown source."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2])

CJK_FONT = "微软雅黑"
LATIN_FONT = "Segoe UI"
BODY_FONT = "宋体"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def set_run_font(run, latin=LATIN_FONT, cjk=CJK_FONT, size=None, bold=None, color=None):
    run.font.name = latin
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)
    rfonts.set(qn("w:cs"), latin)


def style_fonts(style, latin=LATIN_FONT, cjk=CJK_FONT):
    style.font.name = latin
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), cjk)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text, size=Pt(10.5), base_bold=False, color=None):
    text = text.replace("\\", "")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, latin="Consolas", cjk=BODY_FONT, size=size, bold=base_bold, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=base_bold, color=color)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_bottom_border(paragraph, color="1F4E79", size="8"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def build():
    doc = Document()

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    usable = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    style_fonts(normal, latin="Segoe UI", cjk=BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)

    for name, size, color in (("Heading 1", 17, ACCENT), ("Heading 2", 13.5, ACCENT), ("Heading 3", 11.5, RGBColor(0x33, 0x33, 0x33))):
        st = doc.styles[name]
        style_fonts(st)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    # ---- footer: page number ----
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("WorldBranch · 项目总设计文档 v0.1   —   ")
    set_run_font(run, size=Pt(8), color=RGBColor(0x88, 0x88, 0x88))
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer_p._p.append(fld)

    lines = SRC.read_text(encoding="utf-8").splitlines()

    # ---- cover ----
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("WorldBranch")
    set_run_font(r, size=Pt(36), bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("安史之乱模拟器")
    set_run_font(r, size=Pt(22), bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("项目总设计文档")
    set_run_font(r, size=Pt(15), color=RGBColor(0x44, 0x44, 0x44))
    for _ in range(6):
        doc.add_paragraph()

    meta = [
        ("文档版本", "v0.1（设计草案）"),
        ("日期", "2026-10-27"),
        ("状态", "核心方向已确定，部分模块待定"),
        ("适用对象", "单人开发者 + AI 协作"),
        ("文档定位", "项目后续开发的基础依据，不是最终方案"),
    ]
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta:
        cells = tbl.add_row().cells
        cells[0].width = Cm(4.0)
        cells[1].width = usable - Cm(4.0)
        pk = cells[0].paragraphs[0]
        rk = pk.add_run(k)
        set_run_font(rk, size=Pt(10), bold=True, color=ACCENT)
        shade(cells[0], "EAF1F8")
        pv = cells[1].paragraphs[0]
        rv = pv.add_run(v)
        set_run_font(rv, size=Pt(10))
    doc.add_page_break()

    # ---- body ----
    i = 0
    skipped_first_h1 = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---":
            i += 1
            continue

        # table block
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(b) for b in block if not is_separator(b)]
            if not rows:
                continue
            ncol = max(len(r_) for r_ in rows)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            col_w = usable // ncol
            for ridx, cells_txt in enumerate(rows):
                cells = t.add_row().cells
                for cidx in range(ncol):
                    txt = cells_txt[cidx] if cidx < len(cells_txt) else ""
                    cell = cells[cidx]
                    cell.width = col_w
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(2)
                    para.paragraph_format.line_spacing = 1.2
                    if ridx == 0:
                        add_inline(para, txt, size=Pt(9.5), base_bold=True, color=ACCENT)
                        shade(cell, "EAF1F8")
                    else:
                        add_inline(para, txt, size=Pt(9.5))
                        if ridx % 2 == 0:
                            shade(cell, "F7F9FB")
            repeat_header(t.rows[0])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # headings
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not skipped_first_h1:
                skipped_first_h1 = True
                i += 1
                continue
            # level 0 = Title; chapter headings use "Heading 1" so the
            # outline levels stay 1/2/3, deeper markdown levels map upward.
            h = doc.add_heading(level=max(1, level - 1))
            add_inline(h, text, size=Pt({1: 17, 2: 13.5, 3: 11.5, 4: 11}[level]),
                       base_bold=True,
                       color=ACCENT if level <= 2 else RGBColor(0x33, 0x33, 0x33))
            for hr in h.runs:
                hr.font.bold = True
            if level == 1:
                add_bottom_border(h)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip()[1:].strip())
                i += 1
            q = doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.6)
            q.paragraph_format.space_before = Pt(4)
            q.paragraph_format.space_after = Pt(6)
            add_inline(q, " ".join(x for x in buf if x), size=Pt(9.5), color=RGBColor(0x44, 0x44, 0x44))
            add_bottom_border(q, color="BFCEDD", size="6")
            i += 1
            continue

        # ordered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.75)
            add_inline(p, m.group(2))
            i += 1
            continue

        # bullet list
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.75)
            add_inline(p, stripped[2:])
            i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1

    doc.core_properties.title = "WorldBranch 项目总设计文档"
    doc.core_properties.subject = "安史之乱模拟器"
    doc.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    build()